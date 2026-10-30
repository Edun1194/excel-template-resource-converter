#!/usr/bin/env python3
"""PANICALE FW26 STOCK LIST 图片匹配 — 货号+色号合并在 c1 的库存清单版式。

版式(2026-10-30 新增,与 8/28 的 GIACENZE/GIANCE 版式不同):
- 单 sheet, R1 表头(中文):
    c1 = "商品国际编码 + 色号" / "国际编码"  → 两种行:
        · 单 token (如 "U390010PA")            = 尺码行, c4-c13 = 尺码标签
        · 三 token (如 "U390010PA 39U36 0200")= 颜色行, c4-c13 = 各尺码数量
    c2 = 商品描述, c3 = 颜色, c4-c13 = 尺码矩阵,
    c14 = 材质, c15 = 数量, c16 = cost出厂(欧元), c17 = 加价率, c18 = 供货价, c19 = 原价

图片匹配:
- 图库目录 `<V1067>/PANICALE高清图-货号+色号/FW26 {MEN|WOMEN}/`
- 文件名 `<SKU> <色号> <色号数字>.<ext>`, 与 c1 一一对应
- 精确匹配: 归一化(空白折叠 / → - / 大写)后按整串查
- 兜底(仅色号前缀漏写 '3',sku 与色号数字完全一致才算):
  图库 "9U30A 0591" ↔ 清单 "39U30A 0591"  → 命中并单独计数
- 不放宽其他变体: miss 就在报告里列出来,不做颜色/款式兜底

嵌入(保清晰度):
- 原图 bytes 一字不动,只用 OneCellAnchor 改 ext cx/cy → Excel 按 92pt 渲染
- 输出为新文件(源文件不动),追加 "图片" 列

用法:
    python3 match_panicale_stocklist_images.py \
        --men-src "配置器测试文件/V1067/PANICALE FW26 STOCK LIST MEN 30.10.xlsx" \
        --women-src "配置器测试文件/V1067/PANICALE FW26 STOCK LIST WOMEN 30.10.xlsx"
"""

from __future__ import annotations

import argparse
import os
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------
V1067_BASE = Path("/Users/superlist/Desktop/表格处理/配置器测试文件/V1067")
IMG_BASE_DEFAULT = V1067_BASE / "PANICALE高清图-货号+色号"
OUT_DIR_DEFAULT = V1067_BASE / "converted"

IMG_COL_WIDTH = 92.0   # 图片列列宽(字符)
IMG_ROW_HEIGHT = 92.0  # 有图行行高(pt)
IMG_MAX_PT = 92.0      # 图片最长边(pt)
PT_TO_EMU = 12700
PX_TO_EMU = 9525

SUPPORTED_EXTS = {".jpg", ".jpeg", ".png", ".webp"}

# 清单里 c1 的两种形态
SIZE_ROW_TOKENS = 1
COLOR_ROW_TOKENS = 3


# ---------------------------------------------------------------------------
# 图库索引
# ---------------------------------------------------------------------------
def normalize_key(value: str) -> str:
    """归一化: '/' → '-', 折叠空白, 大写。

    "D390010PA/ND 39D49B 0935" → "D390010PA-ND 39D49B 0935"
    "U390864GB 39U31A  0285"  → "U390864GB 39U31A 0285"(双空格折叠)
    """
    return " ".join(str(value).replace("/", "-").split()).upper()


def build_image_index(img_dir: Path) -> Dict[str, Path]:
    """扫目录 → {归一化 stem: 图片路径}。同 key 多张取第一个。"""
    index: Dict[str, Path] = {}
    if not img_dir.is_dir():
        raise FileNotFoundError(f"图片目录不存在: {img_dir}")
    for name in sorted(os.listdir(img_dir)):
        ext = os.path.splitext(name)[1].lower()
        if ext not in SUPPORTED_EXTS:
            continue
        stem = os.path.splitext(name)[0]
        index.setdefault(normalize_key(stem), img_dir / name)
    return index


def find_image(index: Dict[str, Path], raw_c1: str) -> Tuple[Optional[Path], str]:
    """按 c1 找图。返回 (图片路径, 命中方式) / (None, "miss")。

    命中方式: "exact" = 文件名直接对上; "typo" = 图库色号漏写前导 '3'。
    """
    key = normalize_key(raw_c1)
    if key in index:
        return index[key], "exact"

    # 兜底: 色号前缀漏写 '3' (39U30A → 9U30A),sku 与色号数字必须完全一致
    parts = key.split(" ")
    if len(parts) == 3 and parts[1].startswith("3"):
        alt = normalize_key(f"{parts[0]} {parts[1][1:]} {parts[2]}")
        if alt in index:
            return index[alt], "typo"
    return None, "miss"


# ---------------------------------------------------------------------------
# 嵌入
# ---------------------------------------------------------------------------
def load_image_keep_bytes(path: Path) -> Optional[XLImage]:
    """读原图,bytes 不动(保清晰度)。"""
    try:
        return XLImage(str(path))
    except Exception as exc:  # noqa: BLE001
        print(f"  WARN: 无法加载图片 {path.name}: {exc}")
        return None


def compute_display_emu(w_px: int, h_px: int, max_pt: float = IMG_MAX_PT) -> Tuple[int, int]:
    """按 aspect 缩到最长边 ≤ max_pt,返回 (cx, cy) EMU。"""
    max_emu = int(max_pt * PT_TO_EMU)
    cx = w_px * PX_TO_EMU
    cy = h_px * PX_TO_EMU
    if max(cx, cy) <= max_emu:
        return cx, cy
    if cx > cy:
        return max_emu, int(cy * max_emu / cx)
    return int(cx * max_emu / cy), max_emu


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def process_file(
    src_xlsx: Path,
    img_dir: Path,
    out_xlsx: Path,
    label: str,
) -> Dict[str, object]:
    print(f"\n--- {label} ---")
    print(f"src: {src_xlsx.name}")
    print(f"img: {img_dir}")

    index = build_image_index(img_dir)
    print(f"  图库索引: {len(index)} 个文件")

    wb = openpyxl.load_workbook(src_xlsx)
    ws = wb.worksheets[0]
    sheet_name = ws.title

    img_col = ws.max_column + 1
    ws.cell(1, img_col, value="图片")
    ws.column_dimensions[get_column_letter(img_col)].width = IMG_COL_WIDTH

    hit_exact = 0
    hit_typo = 0
    miss_rows: List[Tuple[int, str]] = []
    size_rows = 0
    embedded_bytes = 0
    used_files: Dict[str, str] = {}

    for r in range(2, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        if c1 is None or not str(c1).strip():
            continue
        parts = str(c1).strip().split()
        if len(parts) == SIZE_ROW_TOKENS:
            size_rows += 1
            continue
        if len(parts) != COLOR_ROW_TOKENS:
            continue

        path, how = find_image(index, str(c1))
        if path is None:
            miss_rows.append((r, str(c1).strip()))
            continue

        img = load_image_keep_bytes(path)
        if img is None:
            miss_rows.append((r, str(c1).strip()))
            continue

        cx, cy = compute_display_emu(img.width, img.height, IMG_MAX_PT)
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=img_col - 1, row=r - 1, colOff=0, rowOff=0),
            ext=XDRPositiveSize2D(cx=cx, cy=cy),
        )
        ws.add_image(img)
        ws.row_dimensions[r].height = IMG_ROW_HEIGHT

        embedded_bytes += path.stat().st_size
        used_files[path.name] = how
        if how == "exact":
            hit_exact += 1
        else:
            hit_typo += 1

    total = hit_exact + hit_typo + len(miss_rows)
    rate = (hit_exact + hit_typo) / total * 100 if total else 0.0
    print(f"  颜色行: {total}   尺码行(无色号, 不配图): {size_rows}")
    print(f"  命中: {hit_exact + hit_typo} / {total}  ({rate:.1f}%)"
          f"   [精确 {hit_exact} / 漏写'3'兜底 {hit_typo}]")
    print(f"  未命中: {len(miss_rows)}")
    print(f"  图片列: 第 {img_col} 列 ({get_column_letter(img_col)})")
    print(f"  嵌入图总大小: {embedded_bytes / 1024 / 1024:.1f} MB")

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    print(f"  ✓ saved: {out_xlsx}")

    return {
        "label": label,
        "src": src_xlsx,
        "out": out_xlsx,
        "sheet": sheet_name,
        "img_col": img_col,
        "total": total,
        "hit": hit_exact + hit_typo,
        "hit_exact": hit_exact,
        "hit_typo": hit_typo,
        "miss": miss_rows,
        "index_size": len(index),
        "used_files": used_files,
    }


def write_report(results: List[Dict[str, object]], out_path: Path) -> None:
    """写一份缺图清单(markdown),方便跟供应商要图。"""
    lines = [
        "# PANICALE FW26 STOCK LIST 图片匹配报告",
        "",
        f"生成时间: {__import__('datetime').datetime.now():%Y-%m-%d %H:%M}",
        "",
        "## 汇总",
        "",
        "| 文件 | 颜色行 | 命中 | 精确命中 | 漏写'3'兜底 | 未命中 | 命中率 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for res in results:
        rate = res["hit"] / res["total"] * 100 if res["total"] else 0.0
        lines.append(
            f"| {res['src'].name} | {res['total']} | {res['hit']} | {res['hit_exact']} "
            f"| {res['hit_typo']} | {len(res['miss'])} | {rate:.1f}% |"
        )

    for res in results:
        lines += [
            "",
            f"## {res['label']} — 缺图明细 ({len(res['miss'])} 条)",
            "",
            "| 行号 | 货号 + 色号 |",
            "|---|---|",
        ]
        for r, key in res["miss"]:
            lines.append(f"| {r} | {key} |")

    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"\n报告已写: {out_path}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--men-src", default=str(V1067_BASE / "PANICALE FW26 STOCK LIST MEN 30.10.xlsx"))
    ap.add_argument("--women-src", default=str(V1067_BASE / "PANICALE FW26 STOCK LIST WOMEN 30.10.xlsx"))
    ap.add_argument("--img-base", default=str(IMG_BASE_DEFAULT))
    ap.add_argument("--out-dir", default=str(OUT_DIR_DEFAULT))
    ap.add_argument("--report-name", default="PANICALE_FW26_STOCK_LIST_图片匹配报告.md")
    args = ap.parse_args()

    img_base = Path(args.img_base)
    out_dir = Path(args.out_dir)

    jobs = [
        ("MEN", Path(args.men_src), img_base / "FW26 MEN"),
        ("WOMEN", Path(args.women_src), img_base / "FW26 WOMEN"),
    ]

    results = []
    for label, src, img_dir in jobs:
        if not src.exists():
            print(f"SKIP {label}: {src} 不存在")
            continue
        out_xlsx = out_dir / f"{src.stem}_with_images.xlsx"
        results.append(process_file(src, img_dir, out_xlsx, label))

    if results:
        write_report(results, out_dir / args.report_name)

        print("\n=== SUMMARY ===")
        for res in results:
            rate = res["hit"] / res["total"] * 100 if res["total"] else 0.0
            print(f"  {res['label']}: {res['hit']}/{res['total']} ({rate:.1f}%) "
                  f"→ {res['out'].name}")


if __name__ == "__main__":
    main()
