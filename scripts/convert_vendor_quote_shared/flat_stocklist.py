"""GUCCI 工厂「平铺库存清单」版式解析(V1038 / V1061 共用)。

与 ``col14_parsers.py`` 同理:本模块只负责把文件解析成 ``List[Dict]``,
品牌/仓库 ID 走各 vendor adapter 自己声明(``get_vendor_meta`` + ``VENDOR_ID``)。

适用文件(单 sheet, 意大利语表头):
- ``GUCCI FW26 - WAREHOUSE A.xlsx``  → V1038
- ``GUCCI FW26 - WAREHOUSE B.xlsx``  → V1061

结构(自上而下, 数据行与段落标题行交替):
    R1  : 标题 "GUCCI — Stocklist FW26"          → 品牌 + 季节
    R2  : 空
    R3  : 表头行(Foto | Modello/Variante | Gruppo | Descrizione | Colore
               | MADE IN | Acquisto € | Vendita € | TOTALE | <尺码列>)
    R4+ : 数据行 / 段落标题行交替

段落标题行(只有 c2 有值, c3 空):
    'Donna (38 articoli/varianti)'  → 性别段(女)
    'Uomo  (38 articoli/varianti)'  → 性别段(男)
    'Abbigliamento' / 'Calzature' / 'Accessori' → 类目段(挂在当前性别下)
    'TOTALE PEZZI GUCCI'            → 合计行(跳过)

数据行(c2 货号 和 c3 Gruppo 都有值):
    c2 = Modello/Variante (货号)      c3 = Gruppo    (商品类型)
    c4 = Descrizione    (材质成分)     c5 = Colore    (颜色)
    c6 = MADE IN        (产地)         c7 = Acquisto €(cost)
    c8 = Vendita €      (零售价)       c9 = TOTALE    (总件数)
    c10.. = 尺码矩阵, 表头是尺码标签, 单元格值 = 该尺码库存数

字段映射(源 → 模板, 一一对应不造值):
    Gruppo       → 描述   (DOWN JACKETS / SNEAKERS / TIES …)
    段落大类      → 类目   (Abbigliamento / Calzature / Accessori)
    Descrizione  → 材质   (70% CO 30% PL …)
    标题里的 FW26 → 季节
    (源文件没有色号列) → 色号 留空
    尺码列表头    → 尺码, 该格数值 → 数量

⚠️ 一行一个尺码: 同一货号有多个尺码库存 → 展开成多行(数量取该尺码格的值)。
   图片按 anchor 行号取(c1 'Foto' 列), 同一货号的多个尺码行共用同一张图。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from . import load_images_by_row_per_sheet, to_float


# 固定字段: 归一化后的表头 → 字段名
_FIXED_HEADERS: Dict[str, str] = {
    "foto": "image",
    "modellovariante": "sku",
    "gruppo": "group",
    "descrizione": "material",
    "colore": "color",
    "madein": "origin",
    "acquisto": "cost",
    "vendita": "retail",
    "totale": "total",
}

# 性别段关键词 → 模板性别值(源文件是意大利语)
_GENDER_KEYS: Tuple[Tuple[str, str], ...] = (
    ("donna", "WOMAN"),
    ("uomo", "MAN"),
    ("bambin", "KIDS"),
)

# 季节: FW26 / SS27 / AW27 / CW26 (源标题里)
_SEASON_RE = re.compile(r"\b(FW|SS|AW|CW)\s*[-/]?\s*(\d{2})\b", re.IGNORECASE)

# 半码写法归一: '7/' → '7½'
#   WAREHOUSE B 的 2-11 鞋码段用 ASCII 斜杠写半码('7/' '8/'), 但**同一文件**
#   另一段用的是标准 '½' 字符('36½' '37½' '38½'); WAREHOUSE A 同位置是 '7½'。
#   所以 '/' 在这个版式里就是半码 —— 有源内证据, 不是猜。
_HALF_SLASH_RE = re.compile(r"^(\d+(?:\.\d+)?)/$")

# 源 'Descrizione' 列(c4)的两种语义 —— 同一版式, 不同品牌填的东西不一样:
#   - 材质成分   GUCCI: '70% CO 30% PL'                        → 材质列
#   - 商品描述   CANADA GOOSE: 'CHILLIWACK PUFFER - TD Black - Noir' → 描述列
# 按**内容**判定(整份文件一个模式), 不按品牌名硬编码。
# 判据: 多数数据行的 c4 含 '%' 成分写法 → 材质模式; 否则 → 描述模式。
# 实测: GUCCI A 76/76 行含 %, GUCCI B 45/45 行含 %, CANADA GOOSE 0/7 行含 %。

# 源里的"未填写"占位值(意大利语等)→ 当作没有数据, 输出留空
_NOT_INDICATED = {
    "non indicato", "non indicata", "non specificato", "non specificata",
    "non indicato (n/a)", "n/d", "nd", "n.a.", "n.a", "na", "n.a",
    "-", "--", "---", "—", "/", "null", "none",
}


def _norm_hdr(value: object) -> str:
    """表头归一化: 小写 + 只留字母数字。'Acquisto €' → 'acquisto', 'MADE IN' → 'madein'.

    ⚠️ 只用于**匹配固定字段名**; 尺码标签必须用原值(否则 '37½' 会被抹成 '37')。
    """
    if value is None:
        return ""
    return re.sub(r"[^0-9a-z]+", "", str(value).strip().lower())


def normalize_size_label(label: str) -> str:
    """尺码标签归一: '7/' → '7½'。其余原样返回。"""
    m = _HALF_SLASH_RE.match((label or "").strip())
    if m:
        return f"{m.group(1)}½"
    return (label or "").strip()


def extract_season(text: str) -> str:
    """从标题里抽季节: 'GUCCI — Stocklist FW26' → 'FW26'。取不到就留空。"""
    m = _SEASON_RE.search(text or "")
    if not m:
        return ""
    return f"{m.group(1).upper()}{m.group(2)}"


def extract_brand(text: str) -> str:
    """从标题里抽品牌: 'GUCCI — Stocklist FW26' → 'GUCCI'。取不到就留空。"""
    t = (text or "").strip()
    for sep in ("—", "–", "-", "·", "Stocklist", "stocklist"):
        if sep in t:
            head = t.split(sep, 1)[0].strip()
            if head:
                return head
    return ""


def find_header(
    ws,
) -> Tuple[int, Dict[str, int], List[Tuple[int, str]]]:
    """找平铺库存清单版式的表头行。

    Returns:
        (表头行号, 固定字段列映射 {字段名: 列号}, 尺码列列表 [(列号, 尺码标签)])
        不是这个版式就返回 (0, {}, [])。
    """
    for r in range(1, min(16, ws.max_row) + 1):
        cols: Dict[str, int] = {}
        for c in range(1, min(ws.max_column, 40) + 1):
            key = _norm_hdr(ws.cell(r, c).value)
            if key in _FIXED_HEADERS and key not in cols:
                cols[key] = c
        # 必须同时命中 'Modello/Variante' 和 'TOTALE' —— 才认得是这个版式
        if "modellovariante" not in cols or "totale" not in cols:
            continue
        sizes: List[Tuple[int, str]] = []
        for c in range(cols["totale"] + 1, ws.max_column + 1):
            label = ws.cell(r, c).value
            if label is None or not str(label).strip():
                continue
            sizes.append((c, normalize_size_label(str(label))))  # '7/' → '7½'
        return r, cols, sizes
    return 0, {}, []


def _is_placeholder(value: str) -> bool:
    """源里的"未填写"占位值(Non indicato / N/D / …)→ 当作没有数据。"""
    v = (value or "").strip().lower()
    if not v:
        return True
    return v in _NOT_INDICATED


def _detect_desc_mode(ws, cols: Dict[str, int], hdr_row: int) -> str:
    """判断源 c4 ('Descrizione') 列的语义。

    Returns:
        "material"     → c4 是材质成分(描述取 Gruppo, 类目取段落大类)
        "description"  → c4 是商品描述(描述取 c4, 类目取 Gruppo, 材质留空)
    """
    total = 0
    material = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        sku_v = ws.cell(r, cols["modellovariante"]).value
        sku = str(sku_v).strip() if sku_v is not None else ""
        grp_v = ws.cell(r, cols["gruppo"]).value
        grp = str(grp_v).strip() if grp_v is not None else ""
        if not sku or not grp:
            continue
        total += 1
        v = ws.cell(r, cols["descrizione"]).value
        if v is not None and "%" in str(v):
            material += 1
    if total and material * 2 >= total:
        return "material"
    return "description"


def parse_rows(
    primary_file: Path, vendor_id: str
) -> Optional[List[Dict[str, object]]]:
    """解析平铺库存清单版式。

    不是这个版式就返回 ``None``(让各 vendor adapter 走自己的老路径)。
    """
    rows: List[Dict[str, object]] = []
    wb = openpyxl.load_workbook(primary_file, data_only=True)
    try:
        # 找命中的 sheet + 表头
        target = None
        for name in wb.sheetnames:
            ws = wb[name]
            hdr_row, cols, sizes = find_header(ws)
            if hdr_row:
                target = (name, ws, hdr_row, cols, sizes)
                break
        if target is None:
            return None
        sheet_name, ws, hdr_row, cols, sizes = target

        # 图片: 按 anchor 行号(0-idx)取, 数据行 r(1-idx) → r-1
        images_by_sheet = load_images_by_row_per_sheet(primary_file)
        images_by_row = images_by_sheet.get(sheet_name, {})

        # 标题行 → 品牌 + 季节
        title = ""
        for r in range(1, hdr_row):
            v = ws.cell(r, 1).value
            if v and str(v).strip():
                title = str(v).strip()
                break
        season = extract_season(title)
        brand = extract_brand(title) or sheet_name.strip()

        gender = ""
        category = ""
        desc_mode = _detect_desc_mode(ws, cols, hdr_row)
        for r in range(hdr_row + 1, ws.max_row + 1):
            sku_v = ws.cell(r, cols["modellovariante"]).value
            sku = str(sku_v).strip() if sku_v is not None else ""
            if not sku:
                continue
            grp_v = ws.cell(r, cols["gruppo"]).value
            grp = str(grp_v).strip() if grp_v is not None else ""

            if not grp:
                # ---- 段落标题行(只有 c2 有值) ----
                low = sku.lower()
                hit = False
                for kw, g in _GENDER_KEYS:
                    if kw in low:
                        gender = g
                        hit = True
                        break
                if hit:
                    continue
                if low.startswith("totale"):     # 'TOTALE PEZZI GUCCI' 合计行
                    continue
                category = sku               # Abbigliamento / Calzature / Accessori
                continue

            # ---- 数据行 ----
            desc_v = ws.cell(r, cols["descrizione"]).value
            desc = str(desc_v).strip() if desc_v is not None else ""
            col_v = ws.cell(r, cols["colore"]).value
            color = str(col_v).strip() if col_v is not None else ""
            org_v = ws.cell(r, cols["madein"]).value
            origin = str(org_v).strip() if org_v is not None else ""
            cost = to_float(ws.cell(r, cols["acquisto"]).value)
            retail = to_float(ws.cell(r, cols["vendita"]).value)
            image_items = images_by_row.get(r - 1, [])

            # 颜色列被源文件填成了描述(整串重复)→ 这一行没有独立颜色值, 留空。
            # CANADA GOOSE 3/7 行是这种情况(源把整段描述复制进了颜色列)。
            if color and desc and color == desc:
                color = ""

            # 'Non indicato' 等占位值 = 产地没填 → 留空
            if _is_placeholder(origin):
                origin = ""

            # 描述 / 类目 / 材质 —— 取决于 c4 是材质成分还是商品描述
            if desc_mode == "material":
                description = grp          # Gruppo = 商品类型
                row_category = category     # 段落大类(Abbigliamento / Calzature …)
                material = desc            # Descrizione = 材质成分
            else:
                description = desc          # Descrizione = 商品描述
                row_category = grp          # Gruppo = 商品类型(更具体, 筛选更有用)
                material = ""              # 源文件没有材质信息

            for c, label in sizes:
                qty = to_float(ws.cell(r, c).value)
                if qty is None or qty <= 0:
                    continue
                rows.append({
                    "brand": brand,
                    "sku": sku,
                    "image": "",
                    "description": description,
                    "category": row_category,
                    "material": material,
                    "color": color,
                    "color_code": "",           # 源文件没有色号列
                    "gender": gender,
                    "season": season,
                    "origin": origin,
                    "cost": cost,
                    "cplus_text": "",           # 由 dispatcher 的 --cplus / 规则填
                    "supply_price": None,       # 同上
                    "retail_price": retail,
                    "size": label,
                    "quantity": int(round(qty)),
                    "vendor_id": vendor_id,
                    "image_items": image_items,
                })
        return rows
    finally:
        try:
            wb.close()
        except Exception:
            pass
