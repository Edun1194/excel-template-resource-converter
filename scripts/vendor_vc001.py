"""VC001 CANADA GOOSE FW26 ASA vendor adapter (2026-08-06 新建)

适配文件:CANADA GOOSE 加拿大工厂的库存清单 (ASA = 客户/经销商),来自 xlsx。
版式特点 — 横排尺码 + 每产品 2 行(产品行 + 库存行),中间 2 空行:

  R1 = header
  R2-R4 = 空 buffer
  R5+ = 产品块,每 4 行一组:
    R+0: 产品行 — c2 爱莎码(条码,13位数字)/ c3 CN NAME(描述)/
          c4 ARTICLE(款号+色号)/ c5 系列(Style Name)/
          c6-c12 尺码 label(XS/S/M/L/XL/XXL/XXXL)/
          c13 PCS 总库存/ c14 出厂价(cost)/ c15 批发价加币(retail)
    R+1: 库存行 — c6-c12 各尺码数(整数,空 cell = 0)
    R+2-R+3: 空 buffer
  R+4: 下一个产品块起点

字段映射:
  SKU        ← ARTICLE (c4) 拆款号+色号 → 拼成 "<款号>-<色号>"
              e.g. "2081MB  9061 " → "2081MB-9061"
                   "2267WB  9433 " → "2267WB-9433"
  描述       ← CN NAME (c3) e.g. "CANADA GOOSE 黑色羽绒服马甲"
  颜色       ← 留空(色号就是色码,色名描述里没有独立字段)
  色号       ← ARTICLE 拆出的后半段(9061/9433/9467/...)
  风格       ← 系列 (c5) e.g. "GARSON VEST" / "LODGE HOODY"
  季节       ← "" 留空(main 走文件名兜底 → "FW26")
  cost       ← 出厂价 (c14)
  retail     ← 批发价加币 (c15,加币)
  云仓        ← VENDOR_ID = "VC001"
  品牌       ← "" 留空(main 走文件名兜底 → "CANADA GOOSE")
  性别       ← 从 CN NAME 提取 "男" → "Man" / "女" → "Woman" / 否则 ""
  类目       ← 从 系列 猜:VEST → "Vest" / JACKET → "Jacket" / HOODY → "Hoody" / ...

图片:openpyxl ws._images 按 anchor._from.row (0-idx) 分桶,
     每产品 1 张图(实测 15 个产品块,15 张图)。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from convert_vendor_quote_shared import (
    is_positive_qty,
    normalize_text,
    to_float,
)


# 仓库 ID 强制为 VC001(2026-08-06 用户要求)
VENDOR_ID = "VC001"

# R1 是 header (1-idx)
HEADER_ROW = 1


def get_vendor_meta() -> Dict[str, object]:
    """VC001 vendor 元数据。"""
    return {
        "default_image_col_width": 40.0,
    }


# 0-idx 关键列位置(基于 CANADA GOOSE 15 列版式)
COL_BARCODE = 1      # c2 爱莎码(13 位数字,识别产品行用)
COL_CNNAME = 2       # c3 CN NAME(中文描述)
COL_ARTICLE = 3      # c4 ARTICLE(款号+色号)
COL_STYLE = 4        # c5 系列(Style Name)
COL_SIZE_START = 5   # c6 起尺码 label 列(0-idx 5+)
COL_PCS = 12         # c13 PCS 总库存
COL_COST = 13        # c14 出厂价
COL_RETAIL = 14      # c15 批发价加币

# 性别关键词(从 CN NAME 提取)
_GENDER_KEYWORDS = [
    ("男", "Man"),
    ("女", "Woman"),
    ("童", "Kid"),
    ("中性", "Unisex"),
]

# 类目关键词(从 系列 猜,取第一个匹配的)
_CATEGORY_KEYWORDS = [
    ("VEST", "Vest"),
    ("JACKET", "Jacket"),
    ("HOODY", "Hoody"),
    ("BOMBER", "Bomber"),
    ("COAT", "Coat"),
    ("PARKA", "Parka"),
    ("SHELL", "Shell"),
    ("FLEECE", "Fleece"),
    ("PANT", "Pants"),
    ("SWEATER", "Sweater"),
    ("MITTEN", "Mittens"),
    ("HAT", "Hat"),
    ("SCARF", "Scarf"),
    ("BIB", "Bibs"),
    ("SUIT", "Suit"),
]


def _load_images_by_row(ws) -> Dict[int, List[Dict]]:
    """按 openpyxl 0-idx row 分桶图(每行 1 张)。"""
    out: Dict[int, List[Dict]] = {}
    for img in ws._images:
        anchor = img.anchor
        if not hasattr(anchor, "_from") or anchor._from is None:
            continue
        try:
            data = img._data() if hasattr(img, "_data") else None
        except Exception:
            data = None
        if data is None:
            continue
        row = anchor._from.row
        out.setdefault(row, []).append({
            "bytes": data,
            "format": img.format or "jpeg",
        })
    return out


def _build_size_labels_for_row(ws, r: int) -> Dict[int, str]:
    """从指定行 (1-idx r) c6+ 提取尺码 label,跳过 None/空 cell。

    CANADA GOOSE 文件尺码 label 在**每个产品行自己带**(R+0 c6-c12),
    不是全局 R1 header。所以按行读。
    """
    out: Dict[int, str] = {}
    for c in range(COL_SIZE_START, ws.max_column):
        v = ws.cell(row=r, column=c + 1).value
        if v is None:
            if out:
                break  # 已有 token 后面空 → 本行结束
            continue
        sv = str(v).strip()
        if not sv:
            if out:
                break
            continue
        out[c] = sv
    return out


def _split_article(article_raw: str) -> Tuple[str, str]:
    """拆 ARTICLE 字符串 → (style_code, color_code)

    e.g. "2081MB  9061 " → ("2081MB", "9061")
         "2267WB  9433 " → ("2267WB", "9433")
         "5078MB 9061"   → ("5078MB", "9061")
         "2081MB"        → ("2081MB", "")  (没有色号部分)

    规则:按空格拆,前半 = 款号(字母+数字),后半 = 色号(纯数字)
    """
    if not article_raw:
        return "", ""
    sv = str(article_raw).strip()
    if not sv:
        return "", ""
    parts = sv.split()
    if len(parts) >= 2:
        # 前半(字母+数字) = 款号,后半(纯数字或字母+数字) = 色号
        return parts[0].strip(), parts[-1].strip()
    if len(parts) == 1:
        return parts[0].strip(), ""
    return "", ""


def _pick_gender_from_description(cn_name: str) -> str:
    """从 CN NAME 提取性别。

    e.g. "CANADA GOOSE 男士黑色羽绒服" → "Man"
         "CANADA GOOSE 女士白色羽绒服" → "Woman"
         "CANADA GOOSE 黑色羽绒服马甲" → "" (无性别词)
    """
    if not cn_name:
        return ""
    for kw, label in _GENDER_KEYWORDS:
        if kw in cn_name:
            return label
    return ""


def _guess_category_from_style(style_name: str) -> str:
    """从 系列 (Style Name) 猜类目。

    e.g. "GARSON VEST"    → "Vest"
         "LODGE HOODY"    → "Hoody"
         "CYPRESS JACKET" → "Jacket"
         "CHILLIWACK BOMBER" → "Bomber"
         "ALBERNI REVERSIBLE FLEECE BOMB" → "Fleece" (BOMB 是 BOMBER 缩写)
    """
    if not style_name:
        return ""
    sv = style_name.upper()
    for kw, label in _CATEGORY_KEYWORDS:
        if kw in sv:
            return label
    return ""


def _parse_sheet(ws, vendor_id: str) -> List[Dict[str, object]]:
    """解析单个 sheet,返回 row dict 列表。

    横排尺码模式:每个产品行 R+0 自己带尺码 label(c6-c12),
    下一行 R+1 是库存数(允许跳 1-2 空行兜底)。
    R1 是空 header(只有元数据 cell,尺码 label 全空)。
    """
    rows: List[Dict[str, object]] = []

    images_by_row = _load_images_by_row(ws)

    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        # 1) 识别产品行: c2 (0-idx 1) 是 13 位数字(爱莎码) OR c4 (0-idx 3) 非空(ARTICLE)
        c2 = ws.cell(row=r, column=COL_BARCODE + 1).value
        c4 = ws.cell(row=r, column=COL_ARTICLE + 1).value
        is_product_row = (
            (isinstance(c2, str) and c2.strip().isdigit() and len(c2.strip()) >= 10)
            or (isinstance(c4, str) and c4.strip())
        )
        if not is_product_row:
            continue

        # 2) 拆 ARTICLE → (style_code, color_code) → SKU "<款号>-<色号>"
        style_code, color_code = _split_article(c4)
        if not style_code:
            continue
        if color_code:
            sku = f"{style_code}-{color_code}"
        else:
            sku = style_code

        # 3) 该产品行的尺码 label(R+0 c6-c12)
        size_labels = _build_size_labels_for_row(ws, r)
        if not size_labels:
            continue

        # 4) 描述 (CN NAME)
        cn_name = normalize_text(ws.cell(row=r, column=COL_CNNAME + 1).value)
        desc = cn_name  # 直接用中文名

        # 5) 风格 (系列)
        style = normalize_text(ws.cell(row=r, column=COL_STYLE + 1).value)

        # 6) cost / retail
        cost_v = ws.cell(row=r, column=COL_COST + 1).value
        cost = to_float(cost_v)
        retail_v = ws.cell(row=r, column=COL_RETAIL + 1).value
        retail = to_float(retail_v)

        # 7) PCS 总库存(只是参考,实际库存走 R+1 尺码分布)
        pcs_v = ws.cell(row=r, column=COL_PCS + 1).value

        # 8) 性别(从 CN NAME 提取)
        gender = _pick_gender_from_description(cn_name)

        # 9) 类目(从 系列 猜)
        category = _guess_category_from_style(style)

        # 10) 图片(0-idx row = r - 1)
        zero_idx = r - 1
        row_images = images_by_row.get(zero_idx, [])

        # 11) 找库存行:R+1(下一个非空行) c6-c12 尺码数
        #     允许跳 1-2 空行兜底
        qty_row = r + 1
        for offset in (1, 2, 3):
            cand = r + offset
            if cand > ws.max_row:
                break
            cand_c2 = ws.cell(row=cand, column=COL_BARCODE + 1).value
            cand_c4 = ws.cell(row=cand, column=COL_ARTICLE + 1).value
            # 候选行: 没有任何产品行特征 + size_labels 中至少 1 个 cell 是数字
            has_size_num = any(
                isinstance(ws.cell(row=cand, column=c + 1).value, (int, float))
                for c in size_labels.keys()
            )
            if not cand_c2 and not cand_c4 and has_size_num:
                qty_row = cand
                break

        # 12) 按尺码展开(横排:每列 1 个尺码,qty > 0 → 1 行)
        size_count = 0
        for c, size_label in size_labels.items():
            qty_v = ws.cell(row=qty_row, column=c + 1).value
            if not is_positive_qty(qty_v):
                continue
            qty = int(round(float(to_float(qty_v) or 0)))
            if qty <= 0:
                continue
            size_count += 1
            rows.append({
                "brand": "",  # 留空,main 走文件名兜底
                "sku": sku,
                "image": "",
                "description": desc,
                "category": category,
                "material": "",
                "color": "",  # 没独立色名字段
                "color_code": color_code,
                "gender": gender,
                "season": "",  # 留空,main 走文件名兜底
                "origin": "",
                "cost": cost,
                "cplus_text": "",
                "supply_price": None,
                "retail_price": retail,
                "size": size_label,
                "quantity": qty,
                "vendor_id": vendor_id,
                "image_items": row_images,
                "source_row_height": None,
            })

    return rows


def parse_rows(
    primary_file: Path, _: Optional[Path], vendor_id: str,
) -> List[Dict[str, object]]:
    """解析 CANADA GOOSE FW26 ASA .xlsx,返回 row dict 列表。

    仓库 ID 统一用 VENDOR_ID = "VC001" (2026-08-06 用户要求),忽略传入 vendor_id。
    """
    rows: List[Dict[str, object]] = []
    if not primary_file.exists():
        return rows
    # 用模块级 VENDOR_ID 覆盖 vendor_id,保证「云仓」列稳定是 "VC001"
    vendor_id = VENDOR_ID

    wb = openpyxl.load_workbook(primary_file, data_only=True)
    for ws in wb.worksheets:
        # 防御性检查:R1 c2 (0-idx 1) 应是 "爱莎码" 才是 CANADA GOOSE 版式
        h_marker = normalize_text(ws.cell(row=HEADER_ROW, column=COL_BARCODE + 1).value)
        if "爱莎码" not in h_marker:
            continue  # 跳过不认识的 sheet
        rows.extend(_parse_sheet(ws, vendor_id))
        break  # 只处理第一个 sheet
    return rows
